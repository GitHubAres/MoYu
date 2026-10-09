# 墨语 MoYu - Copyright (c) 2026 墨语（MoYu）贡献者
# Licensed under the MIT License. See LICENSE.
"""作品导出引擎：按范围取章节并生成 TXT / DOCX 文件。"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

from .db import DATA_DIR, get_db, word_count

EXPORT_DIR = DATA_DIR / "exports"

_CHAPTER_PAT = re.compile(r"^第[0-9零一二两三四五六七八九十百千万]+[章节回]")
_CATEGORY_LABEL = {
    "character": "人物",
    "location": "地点",
    "item": "物品",
    "worldbuilding": "设定",
    "other": "其他",
}

_DIGITS = "零一二三四五六七八九"


def _cn_num(n: int) -> str:
    """1-9999 转中文数字（够章节序号用）。"""
    if n <= 10:
        return "十" if n == 10 else _DIGITS[n]
    if n < 20:
        return "十" + _DIGITS[n % 10]
    if n < 100:
        s = _DIGITS[n // 10] + "十"
        return s + (_DIGITS[n % 10] if n % 10 else "")
    if n < 1000:
        s = _DIGITS[n // 100] + "百"
        r = n % 100
        if r:
            s += ("零" if r < 10 else "") + _cn_num(r)
        return s
    s = _DIGITS[n // 1000] + "千"
    r = n % 1000
    if r:
        s += ("零" if r < 100 else "") + _cn_num(r)
    return s


def _fmt_title(title: str, options: dict, index: int) -> str:
    """章节标题规范化：已符合「第X章」样式的原样保留，否则补序号前缀。"""
    t = (title or "").strip() or f"第{_cn_num(index)}章"
    if options.get("format_titles") and not _CHAPTER_PAT.match(t):
        return f"第{_cn_num(index)}章 {t}"
    return t


def _paragraphs(content: str) -> list[str]:
    """正文拆段：去掉空行，保留原有段落顺序。"""
    return [p.strip() for p in (content or "").replace("\r\n", "\n").replace("\r", "\n").split("\n") if p.strip()]


def chapters_for_scope(db, work_id: int, scope: dict) -> list[tuple[str, list[dict]]]:
    """按导出范围返回有序 [(卷名, [章节行])]；scope = {"type": all|volumes|chapters, "ids": [...]}。"""
    scope = scope or {}
    stype = scope.get("type", "all")
    ids = [int(i) for i in (scope.get("ids") or [])]
    if stype == "volumes" and ids:
        marks = ",".join("?" * len(ids))
        vols = db.execute(
            f"SELECT * FROM volumes WHERE work_id=? AND id IN ({marks}) ORDER BY sort_order, id",
            [work_id, *ids]).fetchall()
    else:
        vols = db.execute(
            "SELECT * FROM volumes WHERE work_id=? ORDER BY sort_order, id", (work_id,)).fetchall()
    tree: list[tuple[str, list[dict]]] = []
    for v in vols:
        if stype == "chapters" and ids:
            marks = ",".join("?" * len(ids))
            chs = db.execute(
                f"SELECT * FROM chapters WHERE volume_id=? AND id IN ({marks}) ORDER BY sort_order, id",
                [v["id"], *ids]).fetchall()
        else:
            chs = db.execute(
                "SELECT * FROM chapters WHERE volume_id=? ORDER BY sort_order, id", (v["id"],)).fetchall()
        if chs:
            tree.append((v["title"], [dict(c) for c in chs]))
    return tree


def _entities_appendix(work_id: int) -> list[tuple[str, str, str]]:
    """文末附录素材：[(分类名, 名称, 简介)]。"""
    rows = get_db().execute(
        "SELECT category, name, content FROM entities WHERE work_id=? AND archived=0 ORDER BY category, id",
        (work_id,)).fetchall()
    return [(_CATEGORY_LABEL.get(r["category"], r["category"]), r["name"], (r["content"] or "").strip())
            for r in rows]


def export_txt(work: dict, tree: list[tuple[str, list[dict]]], options: dict, out_path: Path):
    """生成 TXT（UTF-8 无 BOM），返回 (文件路径, 字数)。"""
    lines = [f"《{work['title']}》"]
    if work.get("intro"):
        lines += ["", work["intro"].strip()]
    total = 0
    for vol_name, chs in tree:
        lines += ["", "", vol_name]
        for i, ch in enumerate(chs, 1):
            lines += ["", _fmt_title(ch["title"], options, i), ""]
            for para in _paragraphs(ch["content"]):
                lines.append(("　　" + para) if options.get("indent") else para)
            total += word_count(ch["content"])
    if options.get("append_entities"):
        appendix = _entities_appendix(work["id"])
        if appendix:
            lines += ["", "", "附录：设定简介", ""]
            for cat, name, content in appendix:
                lines.append(f"【{cat}】{name}")
                if content:
                    lines.append(content)
                lines.append("")
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return out_path, total


def export_docx(work: dict, tree: list[tuple[str, list[dict]]], options: dict, out_path: Path):
    """生成 DOCX（标题分级 + 正文段落 + 可选首行缩进），返回 (文件路径, 字数)。"""
    doc = Document()
    doc.add_heading(work["title"], level=0)
    if work.get("intro"):
        doc.add_paragraph(work["intro"].strip())
    total = 0
    for vol_name, chs in tree:
        doc.add_heading(vol_name, level=1)
        for i, ch in enumerate(chs, 1):
            doc.add_heading(_fmt_title(ch["title"], options, i), level=2)
            for para in _paragraphs(ch["content"]):
                p = doc.add_paragraph(para)
                if options.get("indent"):
                    p.paragraph_format.first_line_indent = Pt(24)
            total += word_count(ch["content"])
    if options.get("append_entities"):
        appendix = _entities_appendix(work["id"])
        if appendix:
            doc.add_heading("附录：设定简介", level=1)
            for cat, name, content in appendix:
                doc.add_heading(f"【{cat}】{name}", level=2)
                if content:
                    doc.add_paragraph(content)
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path, total

def export_epub(work: dict, tree: list[tuple[str, list[dict]]], options: dict, out_path: Path):
    """
    轻量原生 EPUB 电子书生成器（纯 Python 标准库 zipfile 实现，零额外依赖）。
    包含 mimetype, META-INF/container.xml, content.opf, toc.ncx 及按章节划分的 XHTML 内容。
    返回 (out_path, total_words)。
    """
    import html
    import uuid
    import zipfile
    from xml.sax.saxutils import escape

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    book_id = f"urn:uuid:{uuid.uuid4()}"
    title = html.escape(work.get("title") or "无标题作品")
    author = "墨语创作者"

    manifest_items = []
    spine_items = []
    toc_nav_points = []
    total = 0

    css_content = """@charset "utf-8";
body {
    margin: 5% 8%;
    font-family: "PingFang SC", "Microsoft YaHei", "Source Han Serif SC", serif;
    line-height: 1.8;
    color: #1a1a1a;
}
h1.title {
    text-align: center;
    margin-top: 20%;
    margin-bottom: 2em;
    font-size: 2em;
    font-weight: bold;
}
h1.vol {
    margin-top: 2em;
    font-size: 1.5em;
    color: #1b2a38;
    border-bottom: 1px solid #e5e9ee;
    padding-bottom: 0.3em;
}
h2.chap {
    margin-top: 1.5em;
    font-size: 1.25em;
    color: #2c3e50;
}
p {
    margin: 0.6em 0;
    text-indent: 2em;
}
.appendix-title {
    margin-top: 2em;
    font-size: 1.4em;
    color: #8b0000;
}
.appendix-item {
    margin-bottom: 1.2em;
}
.appendix-cat {
    font-weight: bold;
    color: #333;
}
"""

    chapters_data = []
    ch_idx = 1
    play_order = 1

    for vol_name, chs in tree:
        for ch in chs:
            c_title = _fmt_title(ch.get("title", ""), options, ch_idx)
            c_paras = _paragraphs(ch.get("content", ""))
            c_wc = word_count(ch.get("content", ""))
            total += c_wc
            file_name = f"chapter_{ch_idx:04d}.xhtml"
            chapters_data.append({
                "vol": vol_name,
                "title": c_title,
                "paras": c_paras,
                "file": file_name,
                "id": f"ch_{ch_idx}",
                "play_order": play_order,
            })
            manifest_items.append(f'<item id="ch_{ch_idx}" href="{file_name}" media-type="application/xhtml+xml"/>')
            spine_items.append(f'<itemref idref="ch_{ch_idx}"/>')
            escaped_title = escape(c_title)
            toc_nav_points.append(f'  <navPoint id="np_{ch_idx}" playOrder="{play_order}">\n    <navLabel><text>{escaped_title}</text></navLabel>\n    <content src="{file_name}"/>\n  </navPoint>')
            ch_idx += 1
            play_order += 1

    appendix_data = []
    if options.get("append_entities"):
        appendix = _entities_appendix(work["id"])
        if appendix:
            app_file = "appendix.xhtml"
            appendix_data = appendix
            manifest_items.append(f'<item id="appendix" href="{app_file}" media-type="application/xhtml+xml"/>')
            spine_items.append('<itemref idref="appendix"/>')
            toc_nav_points.append(f'  <navPoint id="np_app" playOrder="{play_order}">\n    <navLabel><text>附录：设定简介</text></navLabel>\n    <content src="{app_file}"/>\n  </navPoint>')

    with zipfile.ZipFile(str(out_path), "w") as zf:
        zf.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)

        container_xml = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>"""
        zf.writestr("META-INF/container.xml", container_xml, compress_type=zipfile.ZIP_DEFLATED)

        zf.writestr("OEBPS/style.css", css_content, compress_type=zipfile.ZIP_DEFLATED)

        toc_map = "\n".join(toc_nav_points)
        toc_ncx = f"""<?xml version="1.0" encoding="UTF-8"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">
  <head>
    <meta name="dtb:uid" content="{book_id}"/>
    <meta name="dtb:depth" content="2"/>
    <meta name="dtb:totalPageCount" content="0"/>
    <meta name="dtb:maxPageNumber" content="0"/>
  </head>
  <docTitle><text>{escape(title)}</text></docTitle>
  <navMap>
{toc_map}
  </navMap>
</ncx>"""
        zf.writestr("OEBPS/toc.ncx", toc_ncx, compress_type=zipfile.ZIP_DEFLATED)

        manifest_str = "\n    ".join(manifest_items)
        spine_str = "\n    ".join(spine_items)
        content_opf = f"""<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" unique-identifier="BookID" version="2.0">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:opf="http://www.idpf.org/2007/opf">
    <dc:title>{escape(title)}</dc:title>
    <dc:creator>{escape(author)}</dc:creator>
    <dc:identifier id="BookID">{book_id}</dc:identifier>
    <dc:language>zh-CN</dc:language>
  </metadata>
  <manifest>
    <item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>
    <item id="style" href="style.css" media-type="text/css"/>
    {manifest_str}
  </manifest>
  <spine toc="ncx">
    {spine_str}
  </spine>
</package>"""
        zf.writestr("OEBPS/content.opf", content_opf, compress_type=zipfile.ZIP_DEFLATED)

        current_vol = None
        for ch in chapters_data:
            vol_header = ""
            if ch["vol"] and ch["vol"] != current_vol:
                current_vol = ch["vol"]
                vol_header = f'<h1 class="vol">{escape(current_vol)}</h1>'
            
            p_tags = "\n".join(f"<p>{escape(p)}</p>" for p in ch["paras"])
            chap_title_esc = escape(ch["title"])
            xhtml = f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html PUBLIC "-//W3C//DTD XHTML 1.1//EN" "http://www.w3.org/TR/xhtml11/DTD/xhtml11.dtd">
<html xmlns="http://www.w3.org/1999/xhtml">
<head>
  <title>{chap_title_esc}</title>
  <link rel="stylesheet" type="text/css" href="style.css"/>
</head>
<body>
  {vol_header}
  <h2 class="chap">{chap_title_esc}</h2>
  {p_tags}
</body>
</html>"""
            zf.writestr(f"OEBPS/{ch['file']}", xhtml, compress_type=zipfile.ZIP_DEFLATED)

        if appendix_data:
            app_items_html = []
            for cat, name, content in appendix_data:
                app_items_html.append(f'''<div class="appendix-item">
  <div class="appendix-cat">【{escape(cat)}】{escape(name)}</div>
  <p>{escape(content)}</p>
</div>''')
            app_items_str = "\n".join(app_items_html)
            app_xhtml = f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html PUBLIC "-//W3C//DTD XHTML 1.1//EN" "http://www.w3.org/TR/xhtml11/DTD/xhtml11.dtd">
<html xmlns="http://www.w3.org/1999/xhtml">
<head>
  <title>附录：设定简介</title>
  <link rel="stylesheet" type="text/css" href="style.css"/>
</head>
<body>
  <h1 class="appendix-title">附录：设定简介</h1>
  {app_items_str}
</body>
</html>"""
            zf.writestr("OEBPS/appendix.xhtml", app_xhtml, compress_type=zipfile.ZIP_DEFLATED)

    return out_path, total

